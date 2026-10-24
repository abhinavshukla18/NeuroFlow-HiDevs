from docx import Document

doc = Document()
doc.sections[0].header.paragraphs[0].text = "NeuroFlow Confidential"
doc.add_heading("Introduction", level=1)
doc.add_paragraph("NeuroFlow is a retrieval augmented generation platform.")
doc.add_heading("Details", level=2)
doc.add_paragraph("It ingests PDFs, Word files, images, CSVs and web pages.")
table = doc.add_table(rows=2, cols=2)
table.rows[0].cells[0].text = "Name"
table.rows[0].cells[1].text = "Score"
table.rows[1].cells[0].text = "Asha"
table.rows[1].cells[1].text = "92"
doc.save("sample.docx")